#!/usr/bin/env python3
"""读取股权表并进行确定性复算。"""

from __future__ import annotations

import argparse
import csv
from dataclasses import asdict, dataclass, field, is_dataclass
from decimal import Decimal, InvalidOperation
import json
import pathlib
import re
import sys


class InputError(ValueError):
    """表示用户输入无法形成有效数字。"""


class UnverifiedError(RuntimeError):
    """表示文件或表格尚不能完成核验。"""


@dataclass(frozen=True)
class HoldingRow:
    name: str
    shares: Decimal | None
    percentage: Decimal | None
    percentage_places: int | None
    line: int
    rights_type: str | None = None
    note: str = ""
    shares_formula: str | None = None
    percentage_formula: str | None = None


@dataclass
class CapTable:
    identifier: str
    source_name: str
    as_of: str | None = None
    denominator_name: str | None = None
    declared_denominator: Decimal | None = None
    rows: list[HoldingRow] = field(default_factory=list)
    total_rows: list[HoldingRow] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class Finding:
    code: str
    severity: str
    location: str
    message: str
    impact: str
    action: str


@dataclass
class AuditResult:
    status: str
    mode_label: str
    summary: str
    findings: list[Finding]
    tables: list[CapTable]
    denominator: Decimal | None = None
    changes: list[dict[str, object]] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


NAME_HEADER = re.compile(r"股东|持有人|权益人")
SHARES_HEADER = re.compile(r"股数|股份数目|持有股份数|持股数量|持股数")
PERCENTAGE_HEADER = re.compile(r"持股比例|百分比|占已发行|比例|[%％]")
RIGHTS_HEADER = re.compile(r"权益类型|权益类别|持股口径")
NOTE_HEADER = re.compile(r"备注|附注|说明")
TOTAL_ROW = re.compile(
    r"合计|总计|小计|^\s*(?:grand\s+total|subtotal|total)\s*[:：]?\s*$",
    re.IGNORECASE,
)
RIGHTS_TERMS = {
    "direct": re.compile(r"直接持股|登记持股"),
    "beneficial": re.compile(r"实益持股|实益权益"),
    "voting": re.compile(r"表决权|投票权"),
}

STATUS_READY = "可以继续使用"
STATUS_REPAIR = "修正后再使用"
STATUS_PENDING = "尚未完成核验"
SEVERITY_REPAIR = "必须修正"
SEVERITY_VERIFY = "需要核实"
SEVERITY_NOTE = "提示"


def parse_number(value: object, unit: Decimal = Decimal("1")) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    text = str(value).strip().replace(",", "").replace("，", "")
    inline_unit = (
        Decimal("100000000")
        if "亿" in text
        else Decimal("10000")
        if "万" in text
        else unit
    )
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if match is None:
        return None
    try:
        return Decimal(match.group(0)) * inline_unit
    except InvalidOperation as exc:
        raise InputError("数字格式无法识别") from exc


def parse_percentage(
    value: object, header: str
) -> tuple[Decimal | None, int | None]:
    if value is None or str(value).strip() == "":
        return None, None
    text = str(value).strip().replace("％", "%")
    match = re.search(r"-?\d+(?:\.\d+)?", text.replace(",", ""))
    if match is None:
        return None, None
    number_text = match.group(0)
    try:
        number = Decimal(number_text)
    except InvalidOperation as exc:
        raise InputError("持股比例格式无法识别") from exc
    raw_places = len(number_text.partition(".")[2])
    percent_scale = "%" in text or "百分比" in header or "%" in header
    if percent_scale or number > 1:
        return number / Decimal("100"), raw_places
    return number, max(0, raw_places - 2)


def _cell_value(cell: object) -> object:
    return cell.get("value") if isinstance(cell, dict) else cell


def _cell_formula(cell: object) -> str | None:
    if isinstance(cell, dict):
        formula = cell.get("formula")
        return str(formula) if formula is not None else None
    return None


def _cell_text(cell: object) -> str:
    value = _cell_value(cell)
    return "" if value is None else str(value).strip()


def _is_blank_row(row: list[object]) -> bool:
    return not any(_cell_text(cell) or _cell_formula(cell) for cell in row)


def _column_indexes(row: list[object]) -> dict[str, int] | None:
    columns: dict[str, int] = {}
    for index, cell in enumerate(row):
        text = _cell_text(cell)
        if "name" not in columns and NAME_HEADER.search(text):
            columns["name"] = index
        if (
            "shares" not in columns
            and SHARES_HEADER.search(text)
            and not PERCENTAGE_HEADER.search(text)
        ):
            columns["shares"] = index
        if "percentage" not in columns and PERCENTAGE_HEADER.search(text):
            columns["percentage"] = index
        if "rights" not in columns and RIGHTS_HEADER.search(text):
            columns["rights"] = index
        if "note" not in columns and NOTE_HEADER.search(text):
            columns["note"] = index
    if {"name", "shares", "percentage"}.issubset(columns):
        return columns
    return None


def _get_cell(row: list[object], index: int | None) -> object:
    if index is None or index >= len(row):
        return None
    return row[index]


def _holding_from_row(
    row: list[object],
    columns: dict[str, int],
    line: int,
    shares_header: str,
    percentage_header: str,
) -> HoldingRow | None:
    name = _cell_text(_get_cell(row, columns["name"]))
    if not name or set(name) <= {"-", ":"}:
        return None
    unit = (
        Decimal("100000000")
        if "亿" in shares_header
        else Decimal("10000")
        if "万" in shares_header
        else Decimal("1")
    )
    shares_cell = _get_cell(row, columns["shares"])
    percentage_cell = _get_cell(row, columns["percentage"])
    percentage, places = parse_percentage(
        _cell_value(percentage_cell), percentage_header
    )
    return HoldingRow(
        name=name,
        shares=parse_number(_cell_value(shares_cell), unit),
        percentage=percentage,
        percentage_places=places,
        line=line,
        rights_type=_cell_text(_get_cell(row, columns.get("rights"))) or None,
        note=_cell_text(_get_cell(row, columns.get("note"))),
        shares_formula=_cell_formula(shares_cell),
        percentage_formula=_cell_formula(percentage_cell),
    )


def _tables_from_matrices(
    matrices: list[tuple[str, list[list[object]]]], source_name: str
) -> list[CapTable]:
    tables: list[CapTable] = []
    for matrix_name, matrix in matrices:
        index = 0
        table_number = 0
        while index < len(matrix):
            columns = _column_indexes(matrix[index])
            if columns is None:
                index += 1
                continue
            table_number += 1
            header = matrix[index]
            shares_header = _cell_text(_get_cell(header, columns["shares"]))
            percentage_header = _cell_text(
                _get_cell(header, columns["percentage"])
            )
            identifier = (
                matrix_name
                if table_number == 1
                else f"{matrix_name}·表{table_number}"
            )
            table = CapTable(identifier=identifier, source_name=source_name)
            index += 1
            while index < len(matrix):
                row = matrix[index]
                if _is_blank_row(row) or _column_indexes(row) is not None:
                    break
                holding = _holding_from_row(
                    row,
                    columns,
                    index + 1,
                    shares_header,
                    percentage_header,
                )
                if holding is not None:
                    target = (
                        table.total_rows if TOTAL_ROW.search(holding.name) else table.rows
                    )
                    target.append(holding)
                index += 1
            if table.rows or table.total_rows:
                tables.append(table)
            if index < len(matrix) and _is_blank_row(matrix[index]):
                index += 1
    return tables


def _extract_csv_tables(path: pathlib.Path) -> list[CapTable]:
    try:
        with path.open(encoding="utf-8-sig", newline="") as stream:
            matrix = [[cell.strip() for cell in row] for row in csv.reader(stream)]
    except (OSError, UnicodeError, csv.Error) as exc:
        raise UnverifiedError("无法读取CSV文件") from exc
    return _tables_from_matrices([(path.stem, matrix)], path.name)


def _extract_text_tables(path: pathlib.Path) -> list[CapTable]:
    try:
        lines = path.read_text(encoding="utf-8").splitlines()
    except (OSError, UnicodeError) as exc:
        raise UnverifiedError("无法读取文本文件") from exc
    matrices: list[tuple[str, list[list[str]]]] = []
    current: list[list[str]] = []
    title = path.stem
    for line in lines + [""]:
        stripped = line.strip()
        if stripped.startswith("#"):
            if current:
                matrices.append((title, current))
                current = []
            title = stripped.lstrip("#").strip() or path.stem
            continue
        if not stripped:
            if current:
                matrices.append((title, current))
                current = []
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        if len(cells) >= 3:
            current.append(cells)
    return _tables_from_matrices(matrices, path.name)


def _extract_xlsx_tables(path: pathlib.Path) -> list[CapTable]:
    try:
        import openpyxl
    except ImportError as exc:
        raise UnverifiedError("读取Excel需要安装工作簿读取组件") from exc
    try:
        formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
        values = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise UnverifiedError("无法读取Excel文件") from exc
    try:
        matrices: list[tuple[str, list[list[object]]]] = []
        for formula_sheet, value_sheet in zip(
            formulas.worksheets, values.worksheets
        ):
            rows: list[list[object]] = []
            for formula_row, value_row in zip(
                formula_sheet.iter_rows(values_only=True),
                value_sheet.iter_rows(values_only=True),
            ):
                rows.append(
                    [
                        {"formula": formula, "value": value}
                        if isinstance(formula, str) and formula.startswith("=")
                        else value
                        for formula, value in zip(formula_row, value_row)
                    ]
                )
            matrices.append((formula_sheet.title, rows))
    finally:
        formulas.close()
        values.close()
    return _tables_from_matrices(matrices, path.name)


def _extract_docx_tables(path: pathlib.Path) -> list[CapTable]:
    try:
        import docx
    except ImportError as exc:
        raise UnverifiedError("读取Word需要安装文档读取组件") from exc
    try:
        document = docx.Document(path)
    except Exception as exc:
        raise UnverifiedError("无法读取Word文件") from exc
    matrices = [
        (
            f"表格{index}",
            [[cell.text.strip() for cell in row.cells] for row in table.rows],
        )
        for index, table in enumerate(document.tables, start=1)
    ]
    return _tables_from_matrices(matrices, path.name)


def extract_tables(path: str | pathlib.Path) -> list[CapTable]:
    source = pathlib.Path(path)
    suffix = source.suffix.lower()
    if suffix == ".csv":
        tables = _extract_csv_tables(source)
    elif suffix in {".md", ".txt"}:
        tables = _extract_text_tables(source)
    elif suffix == ".xlsx":
        tables = _extract_xlsx_tables(source)
    elif suffix == ".docx":
        tables = _extract_docx_tables(source)
    else:
        raise UnverifiedError("不支持的文件类型")
    if not tables:
        raise UnverifiedError("未识别到同时包含股东、股数和持股比例的表格")
    return tables


def percentage_tolerance(places: int | None) -> Decimal:
    digits = 2 if places is None else places
    return Decimal("0.5") * Decimal("1").scaleb(-(digits + 2))


def result_status(findings: list[Finding]) -> str:
    if any(item.severity == SEVERITY_REPAIR for item in findings):
        return STATUS_REPAIR
    if any(item.severity == SEVERITY_VERIFY for item in findings):
        return STATUS_PENDING
    return STATUS_READY


def _decimal_input(value: object | None, label: str) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    parsed = parse_number(value)
    if parsed is None:
        raise InputError(f"{label}无法识别")
    return parsed


def _format_integer(value: Decimal) -> str:
    return f"{value:,.0f}"


def _median(values: list[Decimal]) -> Decimal | None:
    if not values:
        return None
    ordered = sorted(values)
    middle = len(ordered) // 2
    if len(ordered) % 2:
        return ordered[middle]
    return (ordered[middle - 1] + ordered[middle]) / Decimal("2")


def _finding(
    code: str,
    severity: str,
    location: str,
    message: str,
    impact: str,
    action: str,
) -> Finding:
    return Finding(
        code=code,
        severity=severity,
        location=location,
        message=message,
        impact=impact,
        action=action,
    )


def normalized_name(name: str) -> str:
    return re.sub(r"[\s（）()·,，。.、]", "", name).casefold()


def near_equal_shares(left: Decimal, right: Decimal) -> bool:
    if left <= 0 or right <= 0:
        return False
    return abs(left - right) / max(left, right) < Decimal("0.02")


def _rights_category(value: str | None) -> str | None:
    if value is None or not value.strip():
        return None
    for category, pattern in RIGHTS_TERMS.items():
        if pattern.search(value):
            return category
    return normalized_name(value)


def audit_snapshot(
    table: CapTable,
    *,
    denominator_name: str | None = None,
    declared_denominator: object | None = None,
    as_of: str | None = None,
) -> AuditResult:
    denominator = (
        _decimal_input(declared_denominator, "股本分母")
        if declared_denominator is not None
        else table.declared_denominator
    )
    if denominator is not None and denominator <= 0:
        raise InputError("股本分母必须大于零")

    effective_name = denominator_name or table.denominator_name
    effective_as_of = as_of or table.as_of
    findings: list[Finding] = []
    notes: list[str] = []
    complete_rows: list[HoldingRow] = []

    rows_by_name: dict[str, list[HoldingRow]] = {}
    for row in table.rows:
        rows_by_name.setdefault(normalized_name(row.name), []).append(row)
    for same_name_rows in rows_by_name.values():
        if len(same_name_rows) <= 1:
            continue
        display_names = "、".join(row.name for row in same_name_rows)
        findings.append(
            _finding(
                "duplicate_name",
                SEVERITY_REPAIR,
                table.identifier,
                f"同一规范名称重复出现：{display_names}。",
                "同一股东可能被重复计入明细和合计。",
                "合并重复行，或补充说明这些记录为何属于不同权益。",
            )
        )

    for left_index, left in enumerate(table.rows):
        if left.shares is None:
            continue
        for right in table.rows[left_index + 1 :]:
            if right.shares is None:
                continue
            if normalized_name(left.name) == normalized_name(right.name):
                continue
            if near_equal_shares(left.shares, right.shares):
                findings.append(
                    _finding(
                        "possible_nested_disclosure",
                        SEVERITY_VERIFY,
                        table.identifier,
                        (
                            f"疑似重复披露：{left.name}与{right.name}的股数"
                            f"非常接近（{_format_integer(left.shares)}股与"
                            f"{_format_integer(right.shares)}股）。"
                        ),
                        "两行可能存在包含关系，直接相加可能放大持股比例。",
                        "需核对原始附注，确认两行是独立持股还是同一权益的不同披露口径。",
                    )
                )

    rights_categories = {
        category
        for category in (_rights_category(row.rights_type) for row in table.rows)
        if category is not None
    }
    mixed_rights = len(rights_categories) > 1
    if mixed_rights:
        findings.append(
            _finding(
                "mixed_rights_type",
                SEVERITY_VERIFY,
                table.identifier,
                "同一张表混合了直接持股、实益持股或表决权等不同权益口径。",
                "不同权益类别不能直接相加，合计结果不具备同一口径。",
                "按权益类别拆表后分别核验，并回到原始文件确认包含关系。",
            )
        )

    for row in table.rows:
        location = f"{table.identifier}第{row.line}行·{row.name}"
        formula_missing = (
            row.shares_formula is not None and row.shares is None
        ) or (
            row.percentage_formula is not None and row.percentage is None
        )
        if formula_missing:
            findings.append(
                _finding(
                    "formula_result_missing",
                    SEVERITY_VERIFY,
                    location,
                    f"{row.name}包含公式，但文件中没有可读取的计算结果。",
                    "系统无法独立比较公式显示值与复算值。",
                    "先在Excel中完成计算并保存结果，再重新核验。",
                )
            )
            continue
        if row.shares is None or row.percentage is None:
            findings.append(
                _finding(
                    "incomplete_row",
                    SEVERITY_VERIFY,
                    location,
                    f"{row.name}的股数或持股比例缺失。",
                    "无法确认该行是否使用同一股本分母。",
                    "补充缺失数字后重新核验。",
                )
            )
            continue
        if row.shares < 0 or row.percentage < 0 or row.percentage > 1:
            findings.append(
                _finding(
                    "invalid_range",
                    SEVERITY_REPAIR,
                    location,
                    f"{row.name}存在负股数或超出0%至100%的持股比例。",
                    "该行不能作为正常股东持股参与合计。",
                    "确认该行是股东持股还是轧差调整，并修正展示口径。",
                )
            )
            continue
        complete_rows.append(row)

    inferred_bases: list[tuple[HoldingRow, Decimal, Decimal, Decimal]] = []
    if denominator is not None:
        for row in complete_rows:
            expected = row.shares / denominator
            tolerance = percentage_tolerance(row.percentage_places)
            if abs(expected - row.percentage) > tolerance:
                findings.append(
                    _finding(
                        "percentage_mismatch",
                        SEVERITY_REPAIR,
                        f"{table.identifier}第{row.line}行·{row.name}",
                        (
                            f"{row.name}按股本分母{_format_integer(denominator)}股复算为"
                            f"{expected * 100:.4f}%，与表中{row.percentage * 100:.4f}%不一致。"
                        ),
                        "持股比例、稀释或控制权讨论会使用错误数字。",
                        "确认股数、比例和分母是否属于同一时点后修正。",
                    )
                )
    elif not mixed_rights:
        for row in complete_rows:
            if row.shares <= 0 or row.percentage < Decimal("0.01"):
                continue
            tolerance = percentage_tolerance(row.percentage_places)
            lower_percentage = row.percentage - tolerance
            if lower_percentage <= 0:
                continue
            lower_base = row.shares / (row.percentage + tolerance)
            upper_base = row.shares / lower_percentage
            implied = row.shares / row.percentage
            inferred_bases.append((row, lower_base, upper_base, implied))
        if len(inferred_bases) >= 2:
            common_lower = max(item[1] for item in inferred_bases)
            common_upper = min(item[2] for item in inferred_bases)
            if common_lower > common_upper:
                details = "；".join(
                    f"{row.name}隐含{_format_integer(implied)}股"
                    for row, _, _, implied in inferred_bases
                )
                findings.append(
                    _finding(
                        "denominator_mismatch",
                        SEVERITY_REPAIR,
                        table.identifier,
                        f"各行隐含股本分母不一致：{details}。",
                        "不同分母下的比例不能直接相加或比较。",
                        "统一到同一股本时点和同一分母后重新计算。",
                    )
                )
            denominator = _median([item[3] for item in inferred_bases])

    detail_shares = sum(
        (row.shares for row in complete_rows), start=Decimal("0")
    )
    detail_percentage = sum(
        (row.percentage for row in complete_rows), start=Decimal("0")
    )
    if not mixed_rights and detail_percentage > Decimal("1.001"):
        findings.append(
            _finding(
                "percentage_total_over_100",
                SEVERITY_REPAIR,
                table.identifier,
                f"明细持股比例合计为{detail_percentage * 100:.2f}%，明显超过100%。",
                "表内可能存在重复披露、口径混用或错误分母。",
                "逐行核对权益口径和原始附注后修正。",
            )
        )
    elif not mixed_rights and detail_percentage < Decimal("0.999"):
        notes.append(
            f"已披露持股合计{detail_percentage * 100:.2f}%，"
            f"未解释余额{(Decimal('1') - detail_percentage) * 100:.2f}%。"
        )

    for total in [] if mixed_rights else table.total_rows:
        if total.shares is None or total.percentage is None:
            findings.append(
                _finding(
                    "incomplete_row",
                    SEVERITY_VERIFY,
                    f"{table.identifier}第{total.line}行·{total.name}",
                    "总计行的股数或比例缺失。",
                    "无法确认明细是否与总计闭合。",
                    "补充总计行后重新核验。",
                )
            )
            continue
        percentage_gap = abs(total.percentage - detail_percentage)
        if (
            total.shares != detail_shares
            or percentage_gap > percentage_tolerance(total.percentage_places)
        ):
            findings.append(
                _finding(
                    "total_mismatch",
                    SEVERITY_REPAIR,
                    f"{table.identifier}第{total.line}行·{total.name}",
                    (
                        f"总计行写明{_format_integer(total.shares)}股、"
                        f"{total.percentage * 100:.2f}%，明细合计为"
                        f"{_format_integer(detail_shares)}股、"
                        f"{detail_percentage * 100:.2f}%。"
                    ),
                    "总计不能作为后续稀释或交易前后比较的可靠起点。",
                    "修正明细或总计行，并确认是否遗漏其他股东。",
                )
            )

    if effective_name is None:
        notes.append("仅完成算术口径核验，分母性质待确认。")
    else:
        table.denominator_name = effective_name
    if effective_as_of is None:
        notes.append("数据时点未注明，使用前应补充确认。")
    else:
        table.as_of = effective_as_of
    if declared_denominator is not None:
        table.declared_denominator = denominator
    status = result_status(findings)
    summary = {
        STATUS_READY: "表内股数、比例和合计在当前口径下能够闭合。",
        STATUS_REPAIR: "表内存在必须修正的数字矛盾，修正前不宜继续使用。",
        STATUS_PENDING: "仍有关键数据或口径待确认，尚不能完成核验。",
    }[status]
    return AuditResult(
        status=status,
        mode_label="单表体检",
        summary=summary,
        findings=findings,
        tables=[table],
        denominator=denominator,
        notes=notes,
    )


def _complete_row_map(
    table: CapTable, alias_targets: dict[str, str] | None = None
) -> dict[str, HoldingRow]:
    mapping: dict[str, HoldingRow] = {}
    for row in table.rows:
        if row.shares is None or row.percentage is None:
            continue
        key = normalized_name(row.name)
        if alias_targets is not None:
            key = alias_targets.get(key, key)
        mapping[key] = row
    return mapping


def _same_displayed_holding(left: HoldingRow, right: HoldingRow) -> bool:
    if left.shares is None or right.shares is None:
        return False
    if left.percentage is None or right.percentage is None:
        return False
    share_match = left.shares == right.shares or near_equal_shares(
        left.shares, right.shares
    )
    tolerance = max(
        percentage_tolerance(left.percentage_places),
        percentage_tolerance(right.percentage_places),
    )
    return share_match and abs(left.percentage - right.percentage) <= tolerance


def _change_label(
    before_row: HoldingRow | None, after_row: HoldingRow | None
) -> str:
    if before_row is None:
        return "交易后新增出现"
    if after_row is None:
        return "交易后不再出现"
    if after_row.shares > before_row.shares:
        return "股数增加"
    if after_row.shares < before_row.shares:
        return "股数减少"
    if after_row.percentage < before_row.percentage:
        return "股数不变，持股比例被稀释"
    if after_row.percentage > before_row.percentage:
        return "股数不变，持股比例上升"
    return "持股不变"


def compare_tables(
    before: CapTable,
    after: CapTable,
    *,
    before_denominator: object | None = None,
    after_denominator: object | None = None,
    issued_shares: object | None = None,
    cancelled_shares: object | None = None,
    aliases: dict[str, str] | None = None,
) -> AuditResult:
    before_audit = audit_snapshot(
        before, declared_denominator=before_denominator
    )
    after_audit = audit_snapshot(after, declared_denominator=after_denominator)
    findings = [*before_audit.findings, *after_audit.findings]
    notes = [*before_audit.notes, *after_audit.notes]

    before_total = before_audit.denominator
    after_total = after_audit.denominator
    issued = _decimal_input(issued_shares, "新发行股份")
    cancelled = _decimal_input(cancelled_shares, "回购或注销股份")
    if issued is not None and issued < 0:
        raise InputError("新发行股份不得小于零")
    if cancelled is not None and cancelled < 0:
        raise InputError("回购或注销股份不得小于零")

    bridge_provided = issued_shares is not None or cancelled_shares is not None
    if before_total is not None and after_total is not None:
        if bridge_provided:
            issued = issued or Decimal("0")
            cancelled = cancelled or Decimal("0")
            expected_after = before_total + issued - cancelled
            if expected_after != after_total:
                findings.append(
                    _finding(
                        "share_bridge_mismatch",
                        SEVERITY_REPAIR,
                        f"{before.identifier}→{after.identifier}",
                        (
                            f"按交易前总股本{_format_integer(before_total)}股、新发行"
                            f"{_format_integer(issued)}股和回购或注销"
                            f"{_format_integer(cancelled)}股，交易后应为"
                            f"{_format_integer(expected_after)}股；表中为"
                            f"{_format_integer(after_total)}股。"
                        ),
                        "交易前后总股本不能闭合，稀释结果不可靠。",
                        "核对发行、回购或注销数量以及交易后分母。",
                    )
                )
        elif before_total != after_total:
            findings.append(
                _finding(
                    "unexplained_total_change",
                    SEVERITY_VERIFY,
                    f"{before.identifier}→{after.identifier}",
                    (
                        f"总股本由{_format_integer(before_total)}股变为"
                        f"{_format_integer(after_total)}股，但没有提供发行、回购或注销说明。"
                    ),
                    "无法判断总股本变化来自哪一种股份行为。",
                    "补充新发行、回购或注销股份数量后重新核验。",
                )
            )

    alias_after_to_before: dict[str, str] = {}
    for before_name, after_name in (aliases or {}).items():
        alias_after_to_before[normalized_name(after_name)] = normalized_name(
            before_name
        )
    before_map = _complete_row_map(before)
    after_map = _complete_row_map(after, alias_after_to_before)
    before_only = set(before_map) - set(after_map)
    after_only = set(after_map) - set(before_map)

    suspected_pairs: set[tuple[str, str]] = set()
    for before_key in sorted(before_only):
        for after_key in sorted(after_only):
            if _same_displayed_holding(
                before_map[before_key], after_map[after_key]
            ):
                suspected_pairs.add((before_key, after_key))
                findings.append(
                    _finding(
                        "unmatched_name",
                        SEVERITY_VERIFY,
                        f"{before.identifier}→{after.identifier}",
                        (
                            f"{before_map[before_key].name}与{after_map[after_key].name}"
                            "的股数和比例相同或非常接近，但名称不能自动对应。"
                        ),
                        "系统无法判断这是名称变化，还是一名股东退出、另一名股东新增。",
                        "确认两者是否为同一主体；如是，请提供明确名称对应关系。",
                    )
                )

    changes: list[dict[str, object]] = []
    for key in sorted(set(before_map) | set(after_map)):
        before_row = before_map.get(key)
        after_row = after_map.get(key)
        before_shares = before_row.shares if before_row is not None else None
        after_shares = after_row.shares if after_row is not None else None
        before_percentage = (
            before_row.percentage if before_row is not None else None
        )
        after_percentage = after_row.percentage if after_row is not None else None
        share_change = (
            (after_shares or Decimal("0"))
            - (before_shares or Decimal("0"))
        )
        percentage_change = (
            (after_percentage or Decimal("0"))
            - (before_percentage or Decimal("0"))
        )
        changes.append(
            {
                "name": (
                    after_row.name
                    if after_row is not None
                    else before_row.name
                ),
                "before_shares": before_shares,
                "after_shares": after_shares,
                "share_change": share_change,
                "before_percentage": before_percentage,
                "after_percentage": after_percentage,
                "percentage_change": percentage_change,
                "change_label": _change_label(before_row, after_row),
            }
        )

    status = result_status(findings)
    summary = {
        STATUS_READY: "交易前后股本和各股东持股变化在当前口径下能够闭合。",
        STATUS_REPAIR: "交易前后存在必须修正的数字矛盾，修正前不宜继续使用。",
        STATUS_PENDING: "交易前后仍有变化原因或主体对应关系待确认。",
    }[status]
    return AuditResult(
        status=status,
        mode_label="交易前后对表",
        summary=summary,
        findings=findings,
        tables=[before, after],
        denominator=after_total,
        changes=changes,
        notes=notes,
    )


def _format_decimal(value: Decimal | None) -> str:
    if value is None:
        return "待确认"
    if value == value.to_integral_value():
        return f"{value:,.0f}"
    return f"{value.normalize():,f}"


def _format_percentage(value: Decimal | None) -> str:
    if value is None:
        return "待确认"
    return f"{value * 100:.2f}%"


def _render_findings(lines: list[str], findings: list[Finding]) -> None:
    for severity in (SEVERITY_REPAIR, SEVERITY_VERIFY, SEVERITY_NOTE):
        selected = [item for item in findings if item.severity == severity]
        if not selected:
            continue
        lines.extend([f"## {severity}", ""])
        for index, finding in enumerate(selected, start=1):
            lines.extend(
                [
                    f"### {severity}{index}｜{finding.location}",
                    "",
                    f"- 发现：{finding.message}",
                    f"- 影响：{finding.impact}",
                    f"- 建议：{finding.action}",
                    "",
                ]
            )


def _render_snapshot_details(lines: list[str], result: AuditResult) -> None:
    table = result.tables[0] if result.tables else None
    lines.extend(
        [
            "## 分母与时点",
            "",
            f"- 分母名称：{table.denominator_name if table and table.denominator_name else '待确认'}",
            f"- 股本分母：{_format_decimal(result.denominator)}股",
            f"- 数据时点：{table.as_of if table and table.as_of else '待确认'}",
            "",
            "## 关键复算",
            "",
        ]
    )
    if not result.tables:
        lines.extend(["尚无可复算表格。", ""])
        return
    lines.extend(["| 股东 | 股数 | 表中比例 | 隐含分母 |", "|---|---:|---:|---:|"])
    for current in result.tables:
        for row in current.rows:
            implied = (
                row.shares / row.percentage
                if row.shares is not None
                and row.percentage is not None
                and row.percentage > 0
                else None
            )
            lines.append(
                f"| {row.name} | {_format_decimal(row.shares)} | "
                f"{_format_percentage(row.percentage)} | {_format_decimal(implied)} |"
            )
    lines.append("")


def _render_compare_details(lines: list[str], result: AuditResult) -> None:
    before_total = (
        result.tables[0].declared_denominator if len(result.tables) >= 1 else None
    )
    after_total = (
        result.tables[1].declared_denominator if len(result.tables) >= 2 else result.denominator
    )
    lines.extend(
        [
            "## 总股本变化",
            "",
            f"- 交易前总股本：{_format_decimal(before_total)}股",
            f"- 交易后总股本：{_format_decimal(after_total)}股",
            "",
            "## 股东逐项变化",
            "",
            "| 股东 | 交易前股数 | 交易后股数 | 股数变化 | 比例变化 | 变化说明 |",
            "|---|---:|---:|---:|---:|---|",
        ]
    )
    for change in result.changes:
        lines.append(
            f"| {change['name']} | {_format_decimal(change['before_shares'])} | "
            f"{_format_decimal(change['after_shares'])} | "
            f"{_format_decimal(change['share_change'])} | "
            f"{_format_percentage(change['percentage_change'])} | "
            f"{change['change_label']} |"
        )
    lines.extend(["", "## 稀释与新增持股", ""])
    highlighted = [
        change
        for change in result.changes
        if "稀释" in str(change["change_label"])
        or "新增" in str(change["change_label"])
    ]
    if highlighted:
        for change in highlighted:
            lines.append(f"- {change['name']}：{change['change_label']}。")
    else:
        lines.append("- 本次未识别到稀释或新增持股。")
    lines.extend(["", "## 未解释差额", ""])
    unexplained = [
        item
        for item in result.findings
        if item.code in {"share_bridge_mismatch", "unexplained_total_change"}
    ]
    if unexplained:
        for finding in unexplained:
            lines.append(f"- {finding.message}")
    else:
        lines.append("- 当前输入下没有未解释的总股本差额。")
    lines.append("")


def _next_steps(result: AuditResult) -> list[str]:
    if result.status == STATUS_REPAIR:
        return ["先修正必须修正的问题，再重新运行股权表核验。"]
    if result.status == STATUS_PENDING:
        return ["补齐需要核实的分母、附注或股份变化原因，再重新核验。"]
    if result.mode_label == "交易前后对表":
        return [
            "需要比较其他增资、老股或股份支付方案时，进入「交易结构算盘」。",
            "需要保留本次口径和修正轨迹时，保存到「并购项目档案」。",
        ]
    return [
        "需要复算不同交易安排时，进入「交易结构算盘」。",
        "需要保存本次口径和修正轨迹时，保存到「并购项目档案」。",
    ]


def render_markdown(result: AuditResult) -> str:
    lines = [
        "# 股权表核验",
        "",
        "**每一股都对上，每一次变化都有来路。**",
        "",
        "## 一句话结论",
        "",
        f"**{result.status}**｜{result.summary}",
        "",
        "## 使用建议",
        "",
        {
            STATUS_READY: "当前算术口径可以继续使用，仍应保留来源和时点说明。",
            STATUS_REPAIR: "先修正数字矛盾，不要据此继续计算持股、稀释或控制权。",
            STATUS_PENDING: "先补齐缺失事实或确认口径，不要把未发现问题写成核验通过。",
        }[result.status],
        "",
    ]
    if result.mode_label == "交易前后对表":
        _render_compare_details(lines, result)
    else:
        _render_snapshot_details(lines, result)
    _render_findings(lines, result.findings)
    if result.notes:
        lines.extend(["## 补充说明", ""])
        lines.extend(f"- {note}" for note in dict.fromkeys(result.notes))
        lines.append("")
    lines.extend(["## 下一步", ""])
    lines.extend(f"- {step}" for step in _next_steps(result)[:2])
    lines.extend(
        [
            "",
            "本次结果只表示表内算术和用户明确提供的变化能够闭合，"
            "不等于控制权、并表或监管结论。",
            "",
            "## 反馈",
            "",
            "> 这次核验有没有帮你把股权表对清楚？  ",
            "> 已经对清楚｜还有一处不确定｜结果不对｜提交建议",
        ]
    )
    return "\n".join(lines).rstrip() + "\n"


def _json_ready(value: object) -> object:
    if isinstance(value, Decimal):
        return format(value, "f")
    if is_dataclass(value):
        return _json_ready(asdict(value))
    if isinstance(value, dict):
        return {str(key): _json_ready(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_ready(item) for item in value]
    return value


def result_to_dict(result: AuditResult) -> dict[str, object]:
    payload = _json_ready(result)
    if not isinstance(payload, dict):
        raise TypeError("结构化结果必须是对象")
    return payload


def _combined_snapshot_result(
    tables: list[CapTable],
    *,
    denominator_name: str | None,
    declared_denominator: object | None,
    as_of: str | None,
) -> AuditResult:
    results = [
        audit_snapshot(
            table,
            denominator_name=denominator_name,
            declared_denominator=declared_denominator,
            as_of=as_of,
        )
        for table in tables
    ]
    if len(results) == 1:
        return results[0]
    findings = [finding for result in results for finding in result.findings]
    notes = [note for result in results for note in result.notes]
    status = result_status(findings)
    summary = {
        STATUS_READY: f"文件内{len(results)}张股权表在当前口径下均能闭合。",
        STATUS_REPAIR: f"文件内{len(results)}张股权表中存在必须修正的数字矛盾。",
        STATUS_PENDING: f"文件内{len(results)}张股权表中仍有待确认事项。",
    }[status]
    return AuditResult(
        status=status,
        mode_label="单表体检",
        summary=summary,
        findings=findings,
        tables=tables,
        denominator=None,
        notes=notes,
    )


def _parse_aliases(values: list[str]) -> dict[str, str]:
    aliases: dict[str, str] = {}
    for value in values:
        before, separator, after = value.partition("=")
        if not separator or not before.strip() or not after.strip():
            raise InputError("名称对应关系必须使用“交易前名称=交易后名称”")
        aliases[before.strip()] = after.strip()
    return aliases


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="股权表核验",
        description="核验股本分母、持股比例、合计和交易前后变化。",
    )
    subparsers = parser.add_subparsers(dest="mode", required=True)
    snapshot = subparsers.add_parser("snapshot", help="单表体检")
    snapshot.add_argument("file")
    snapshot.add_argument("--denominator-name")
    snapshot.add_argument("--declared-denominator")
    snapshot.add_argument("--as-of")
    snapshot.add_argument("--json", action="store_true")

    compare = subparsers.add_parser("compare", help="交易前后对表")
    compare.add_argument("--before", required=True)
    compare.add_argument("--after", required=True)
    compare.add_argument("--before-denominator")
    compare.add_argument("--after-denominator")
    compare.add_argument("--issued-shares")
    compare.add_argument("--cancelled-shares")
    compare.add_argument("--alias", action="append", default=[])
    compare.add_argument("--json", action="store_true")
    return parser


def _run_from_args(args: argparse.Namespace) -> AuditResult:
    if args.mode == "snapshot":
        return _combined_snapshot_result(
            extract_tables(args.file),
            denominator_name=args.denominator_name,
            declared_denominator=args.declared_denominator,
            as_of=args.as_of,
        )
    before_tables = extract_tables(args.before)
    after_tables = extract_tables(args.after)
    if len(before_tables) != 1 or len(after_tables) != 1:
        raise UnverifiedError("交易前后文件各应只包含一张明确股权表")
    return compare_tables(
        before_tables[0],
        after_tables[0],
        before_denominator=args.before_denominator,
        after_denominator=args.after_denominator,
        issued_shares=args.issued_shares,
        cancelled_shares=args.cancelled_shares,
        aliases=_parse_aliases(args.alias),
    )


def main(argv: list[str] | None = None) -> int:
    parser = _build_parser()
    args = parser.parse_args(argv)
    try:
        result = _run_from_args(args)
    except (InputError, UnverifiedError, OSError, ValueError) as exc:
        result = AuditResult(
            status=STATUS_PENDING,
            mode_label=(
                "交易前后对表" if getattr(args, "mode", None) == "compare" else "单表体检"
            ),
            summary=f"未完成核验：{exc}",
            findings=[],
            tables=[],
            notes=["请补充可识别的股权表或修正输入后重新核验。"],
        )
    if getattr(args, "json", False):
        print(
            json.dumps(
                result_to_dict(result),
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            )
        )
    else:
        print(render_markdown(result), end="")
    return {
        STATUS_READY: 0,
        STATUS_REPAIR: 1,
        STATUS_PENDING: 2,
    }[result.status]


if __name__ == "__main__":
    sys.exit(main())
